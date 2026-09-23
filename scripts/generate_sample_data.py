"""Generate the sample complaint documents in ``data/``.

All people, companies, emails and phone numbers are fictional. Run with:

    PYTHONPATH=src .venv/bin/python scripts/generate_sample_data.py [output_dir]

The helper functions (``write_pdf``, ``write_docx``, ``write_txt``,
``write_corrupted_pdf``) are reused by the test-suite.
"""

from __future__ import annotations

import os
import random
import sys
from pathlib import Path
from typing import Sequence

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

COMPANY = "BrightWave Electronics & Broadband Ltd."
COMPANY_FOOTER = "BrightWave Customer Care | support@brightwave.example.com | 1800-555-0100"

# A section is (heading, body). Body may be a str (paragraphs split on blank lines)
# or a list of (label, value) rows rendered as a key/value table.
Section = tuple[str, "str | list[tuple[str, str]]"]


# --------------------------------------------------------------------------- #
# Writers (also used by tests)
# --------------------------------------------------------------------------- #
def write_pdf(path: Path, title: str, sections: Sequence[Section]) -> Path:
    """Render a simple form-like PDF with a header, key/value tables and paragraphs."""
    styles = getSampleStyleSheet()
    h_company = ParagraphStyle("company", parent=styles["Title"], fontSize=15, spaceAfter=2)
    h_title = ParagraphStyle("title", parent=styles["Heading2"], textColor=colors.HexColor("#1F4E79"))
    h_section = ParagraphStyle("section", parent=styles["Heading4"], spaceBefore=8)
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=10, leading=13)
    footer = ParagraphStyle("footer", parent=body, fontSize=8, textColor=colors.grey)

    story: list = [Paragraph(COMPANY, h_company), Paragraph(title, h_title), Spacer(1, 4)]
    for heading, content in sections:
        story.append(Paragraph(heading, h_section))
        if isinstance(content, str):
            for para in content.strip().split("\n\n"):
                story.append(Paragraph(para.replace("\n", "<br/>"), body))
                story.append(Spacer(1, 3))
        else:
            table = Table([[Paragraph(f"<b>{k}</b>", body), Paragraph(v, body)] for k, v in content],
                          colWidths=[4.5 * cm, 11.5 * cm])
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF3F8")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(table)
    story += [Spacer(1, 12), Paragraph(COMPANY_FOOTER, footer)]

    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=A4, title=title, author="BrightWave Customer Care",
                      leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm,
                      bottomMargin=1.8 * cm).build(story)
    return path


def write_docx(path: Path, title: str, sections: Sequence[Section]) -> Path:
    """Write a Word document; list-type sections become two-column tables."""
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(COMPANY, level=1)
    doc.add_heading(title, level=2)
    for heading, content in sections:
        doc.add_heading(heading, level=3)
        if isinstance(content, str):
            for para in content.strip().split("\n\n"):
                doc.add_paragraph(para)
        else:
            table = doc.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            for key, value in content:
                cells = table.add_row().cells
                cells[0].text, cells[1].text = key, value
    doc.add_paragraph(COMPANY_FOOTER)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def write_txt(path: Path, content: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content.strip() + "\n", encoding="utf-8")
    return path


def write_corrupted_pdf(path: Path, size: int = 2048, seed: int = 7) -> Path:
    """Write random bytes with a .pdf extension (deliberately unreadable)."""
    rng = random.Random(seed)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n" + bytes(rng.getrandbits(8) for _ in range(size)))
    return path


# --------------------------------------------------------------------------- #
# Sample content
# --------------------------------------------------------------------------- #
COMPLAINT_001: list[Section] = [
    ("Case Details", [
        ("Case Reference", "BW-CMP-2026-001"),
        ("Date Received", "02 September 2026"),
        ("Channel", "Web complaint form"),
        ("Assigned Agent", "Priya Raman (Billing Team)"),
    ]),
    ("Customer Details", [
        ("Customer Name", "Arjun Mehta"),
        ("Email", "arjun.mehta@example.com"),
        ("Phone", "+91 98765 40001"),
        ("Account Number", "BW-ACC-448120"),
        ("Product / Service", "FiberMax 300 Mbps Broadband Plan"),
    ]),
    ("Complaint Description",
     "I have been a FiberMax 300 customer for two years and my monthly charge has always been "
     "Rs. 1,199. My August 2026 invoice (INV-2026-08-44812) shows Rs. 2,398, which is exactly "
     "double. The invoice lists the same plan twice with identical dates. I did not add any new "
     "service or upgrade. The full amount was auto-debited from my bank account on 28 August.\n\n"
     "I would like the duplicate charge of Rs. 1,199 refunded and an assurance that this will not "
     "happen again next month."),
    ("Resolution",
     "05 Sep 2026 - Billing team confirmed a duplicate plan line caused by a system migration "
     "error. A refund of Rs. 1,199 was processed to the customer's original payment method "
     "(refund ref RF-77120) and a corrected invoice was issued. Customer confirmed receipt of the "
     "refund by email on 08 Sep 2026 and was satisfied with the outcome."),
    ("Escalation", "Not required. Resolved at first level."),
    ("Attachments", "1. Invoice INV-2026-08-44812.pdf (copy of disputed August invoice)\n"
                    "2. Bank statement extract showing auto-debit on 28-Aug-2026"),
    ("Case Status", "Resolved - closed on 08 Sep 2026."),
]

COMPLAINT_002: list[Section] = [
    ("Case Details", [
        ("Case Reference", "BW-CMP-2026-002"),
        ("Date Received", "10 September 2026"),
        ("Channel", "Email to support@brightwave.example.com"),
        ("Priority", "High"),
    ]),
    ("Customer Details", [
        ("Customer Name", "Sneha Kulkarni"),
        ("Email", "sneha.kulkarni@example.com"),
        ("Phone", "+91 91234 50002"),
        ("Order Number", "BW-ORD-983311"),
        ("Product", "BrightWave NovaBook 14 Laptop (16GB / 512GB), Serial NB14-7Q2X-5531"),
    ]),
    ("Customer Message",
     "Subject: Defective NovaBook 14 - second failure, I want this escalated\n\n"
     "Dear BrightWave team,\n\n"
     "I bought the NovaBook 14 on 18 August 2026. Within ten days the screen started flickering "
     "and the laptop shuts down randomly even at 80% battery. I raised ticket BW-TKT-55120 and "
     "your service centre replaced the motherboard on 01 September. Since then the same problem "
     "has returned and now the keyboard backlight also fails.\n\n"
     "This is a brand new laptop and I rely on it for work. I have attached photos and a short "
     "video of the flickering screen. I do not want another repair. I am asking for a full "
     "replacement or refund, and I want this complaint escalated to a manager because the "
     "first-level support has not solved it. If I do not hear back within 48 hours I will "
     "approach the consumer forum.\n\n"
     "Regards,\nSneha Kulkarni"),
    ("Agent Notes",
     "11 Sep 2026 - Customer contacted. Replacement request is outside first-level authority. "
     "No resolution offered yet; awaiting product quality team review."),
    ("Attachments", "3 photos of flickering display (IMG_2041.jpg, IMG_2042.jpg, IMG_2043.jpg), "
                    "1 video clip, copy of purchase invoice"),
    ("Case Status", "Open - unresolved. Customer requested escalation to manager."),
]

COMPLAINT_003 = """
BRIGHTWAVE ELECTRONICS & BROADBAND - CASE NOTES
===============================================
Case Ref       : BW-CMP-2026-003
Date Opened    : 14 September 2026
Channel        : Phone call to customer care
Agent          : Rahul Verma (Network Support, L1)

CUSTOMER
Name           : Meera Iyer
Email          : meera.iyer@example.com
Phone          : +91 99887 60003
Service        : HomeNet 100 Mbps Broadband, Connection ID HN-220718
Address        : Flat 4B, Lakeview Residency, Koramangala, Bengaluru

ISSUE SUMMARY
Customer reports repeated broadband outages over the last two weeks. The connection drops
completely three to four times a day, usually for 30 to 90 minutes, mostly in the evening.
She works from home and has missed two client video calls. She says this is the third
time she has called; previous tickets BW-TKT-60112 (03 Sep) and BW-TKT-60398 (08 Sep) were
closed after a remote router reboot, but the problem came back each time.

TROUBLESHOOTING DONE
- Remote line test shows high signal loss on the fibre drop to the building.
- Router firmware updated remotely to v4.2.1; issue persists.
- Customer confirmed no other devices or power issues at her end.

ACTION TAKEN / NEXT STEPS
A field technician visit has been scheduled for 16 September 2026, 10:00-13:00, to inspect
the fibre splice box and replace the drop cable if required. A service credit of 5 days will
be reviewed once the fault is fixed. Customer was polite but frustrated; she asked us to
call before the visit.

ATTACHMENTS
None. No documents or screenshots were provided by the customer.

ESCALATION
Not escalated at this stage. Will be reviewed by L2 if the technician visit does not fix the fault.

STATUS
In Progress - technician visit scheduled.
"""

COMPLAINT_004: list[Section] = [
    ("Case Details", [
        ("Case Reference", "BW-CMP-2026-004"),
        ("Date Received", "12 September 2026"),
        ("Channel", "Live chat, transcribed by agent"),
        ("Assigned To", "Logistics Escalations Desk"),
    ]),
    ("Customer Details", [
        ("Customer Name", "Daniel D'Souza"),
        ("Email", "daniel.dsouza@example.com"),
        ("Phone", "+91 90000 70004"),
        ("Delivery Address", "22 Palm Grove Road, Panaji, Goa 403001"),
    ]),
    ("Order Details", [
        ("Order Number", "BW-ORD-771045"),
        ("Order Date", "25 August 2026"),
        ("Item", "BrightWave Vista 55\" 4K Smart TV + wall-mount kit"),
        ("Order Value", "Rs. 48,990 (prepaid)"),
        ("Promised Delivery", "31 August 2026"),
        ("Courier / AWB", "SwiftShip Logistics / SS-449920817"),
        ("Current Tracking", "In transit - held at Pune hub since 02 Sep 2026"),
    ]),
    ("Complaint Description",
     "The customer ordered a 55-inch TV which was promised for delivery by 31 August. It is now "
     "twelve days late and tracking has shown 'held at hub' for ten days with no update. He has "
     "called three times and was told each time it would arrive in 48 hours. He bought the TV "
     "for a family event on 14 September and is extremely unhappy. He is asking for a firm "
     "delivery date, compensation for the delay, or cancellation with a full refund."),
    ("Action Taken",
     "Agent raised a trace request with SwiftShip on 12 Sep. Because the delay exceeds seven days "
     "past the promised date and the customer has contacted us three times, the case has been "
     "escalated to the Logistics Escalations Desk and the Regional Operations Manager as per "
     "policy. No resolution has been confirmed to the customer yet."),
    ("Attachments", "Screenshot of order confirmation email and courier tracking page shared by customer in chat."),
    ("Case Status", "Escalated"),
]

COMPLAINT_005: list[Section] = [
    ("Case Details", [
        ("Case Reference", "BW-CMP-2026-005"),
        ("Date Received", "15 September 2026"),
        ("Channel", "Email"),
        ("Previous Contacts", "4 (tickets BW-TKT-58801, BW-TKT-59230, BW-TKT-59977, BW-TKT-60544)"),
    ]),
    ("Customer Details", [
        ("Customer Name", "Kavita Singh"),
        ("Email", "kavita.singh@example.com"),
        ("Phone", "Not provided"),
        ("Order Number", "BW-ORD-665210"),
        ("Product", "BrightWave SoundPod Pro Wireless Earbuds"),
    ]),
    ("Customer Message",
     "Hello,\n\nI returned the SoundPod Pro earbuds on 05 August 2026 because the left earbud would "
     "not charge. Your courier picked them up (return AWB RT-3321087) and I received an email on "
     "09 August confirming the return was received and inspected at your warehouse. I was told the "
     "refund of Rs. 7,499 would reach my card within 7 working days.\n\n"
     "It has now been more than five weeks and I have not received any money. I have written to "
     "you four times and each time I get a copy-paste reply asking me to wait. Nobody has given me "
     "a transaction reference. I have attached the return pickup receipt and your own return "
     "confirmation email.\n\n"
     "Please escalate this to someone senior who can actually release my refund. Please reply by "
     "email only.\n\nKavita Singh"),
    ("Agent Notes",
     "15 Sep 2026 - Finance shows refund stuck in 'pending approval' state since 10 Aug. Escalated "
     "to Refunds Supervisor (L2) for manual release. Refund not yet issued."),
    ("Attachments", "Return pickup receipt RT-3321087; return confirmation email dated 09-Aug-2026"),
    ("Case Status", "Escalated - refund pending."),
]

COMPLAINT_006 = """
From: Rohan Kapoor <rohan.kapoor@example.com>
To: feedback@brightwave.example.com
Date: 17 September 2026
Subject: Thank you for the quick installation + a question about plans

Reference: BW-FDB-2026-006

Hi BrightWave team,

I just wanted to say thank you. My new FiberMax 500 connection was installed yesterday at our
new flat in Pune and the whole experience was excellent. The technician, Mr. Anil, arrived on
time, explained everything clearly, set up the mesh Wi-Fi in every room and even helped my
father connect his tablet. Speeds have been consistently above 480 Mbps. Please pass my
appreciation on to him and the scheduling team.

I also have a general question: does the FiberMax 500 plan allow adding the BrightWave TV+
streaming bundle mid-cycle, and would the charge be pro-rated? No rush on this, I am just
planning ahead.

You can reach me on +91 98111 80006 if a call is easier.

Best regards,
Rohan Kapoor

---
Internal note (Customer Care, 18 Sep 2026): Positive feedback shared with field team lead.
Customer's plan query answered by email - TV+ can be added any time and is pro-rated from the
activation date. No complaint raised, no action pending, no attachments. Ticket closed.
"""


def generate_all(out_dir: Path) -> list[Path]:
    """Write all seven sample files to ``out_dir`` and return their paths."""
    out_dir.mkdir(parents=True, exist_ok=True)
    return [
        write_pdf(out_dir / "complaint_001.pdf", "Customer Complaint Form - Billing", COMPLAINT_001),
        write_pdf(out_dir / "complaint_002.pdf", "Customer Complaint - Product Defect", COMPLAINT_002),
        write_txt(out_dir / "complaint_003.txt", COMPLAINT_003),
        write_docx(out_dir / "complaint_004.docx", "Complaint Record - Delayed Delivery", COMPLAINT_004),
        write_pdf(out_dir / "complaint_005.pdf", "Customer Complaint - Refund Not Received", COMPLAINT_005),
        write_txt(out_dir / "complaint_006.txt", COMPLAINT_006),
        write_corrupted_pdf(out_dir / "corrupted_007.pdf"),
    ]


if __name__ == "__main__":
    root = Path(__file__).resolve().parents[1]
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else root / os.getenv("DATA_DIR", "data")
    for p in generate_all(target):
        print(f"wrote {p}")
